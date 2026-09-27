"""
03_features.py - Add the fields the analysis needs.

Input:  data/processed/classes_clean.csv   (from 02b_add_class_log.py)
Output: data/processed/classes_features.csv
        data/processed/classes_features.xlsx

New columns:
  day_of_week, day_num, month, time_label, slot   when and where the class ran
  shift_type          Regular or Picked up (rules below)
  capacity, booked, waitlist, fill_rate, is_full  attendance vs room size
  pay_if_standard_rate        what the class pays on the Standard formula
  hit_floor           attendance too low to beat the $28 minimum

Usage:
    python src/03_features.py
"""
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
IN_FILE = ROOT / "data" / "processed" / "classes_clean.csv"
OUT_DIR = ROOT / "data" / "processed"

from config import (CAPACITY, FLOOR_PAY, MAX_PAY, REGULAR_SCHEDULE,  # noqa: E402
                    SCHEDULE_EXCEPTIONS, STUDIO_REV_SHARE)


def is_regular(row) -> bool:
    d = row["class_date"].date()
    for studio, weekday, times, start, end in REGULAR_SCHEDULE:
        if (row["studio_num"] == studio
                and row["class_date"].dayofweek == weekday
                and row["class_time"] in times
                and (start is None or d >= start)
                and (end is None or d <= end)):
            return True
    return False


def add_time_fields(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["day_of_week"] = df["class_date"].dt.day_name()
    df["day_num"] = df["class_date"].dt.dayofweek
    df["month"] = df["class_date"].dt.strftime("%Y-%m")
    df["class_hour"] = df["class_time"].apply(lambda t: t.hour + t.minute / 60)
    df["time_label"] = df["class_time"].apply(
        lambda t: f"{t.hour % 12 or 12}:{t.minute:02d} {'AM' if t.hour < 12 else 'PM'}")
    df["slot"] = (df["studio"] + " | " + df["day_of_week"].str[:3] + " "
                  + df["time_label"])
    return df


def add_shift_type(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    key = list(zip(df["studio_num"], df["class_date"].dt.date, df["class_time"]))
    df["schedule_note"] = [SCHEDULE_EXCEPTIONS.get(k, "") for k in key]
    regular = df.apply(is_regular, axis=1) | df["schedule_note"].ne("")
    df["shift_type"] = regular.map({True: "Regular", False: "Picked up"})

    found = df["schedule_note"].ne("").sum()
    if found != len(SCHEDULE_EXCEPTIONS):
        raise ValueError(f"Only {found} of {len(SCHEDULE_EXCEPTIONS)} schedule "
                         "exceptions matched a class - check dates and times")
    print(f"Applied {found} schedule exceptions (holiday / swapped shifts)")
    return df


def add_attendance(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["capacity"] = df["studio_num"].map(CAPACITY)
    df["booked"] = df[["total_visits", "capacity"]].min(axis=1)
    df["waitlist"] = (df["total_visits"] - df["capacity"]).clip(lower=0)
    df["fill_rate"] = (df["total_visits"] / df["capacity"]).round(3)
    df["is_full"] = df["total_visits"] >= df["capacity"]
    return df


def add_pay_if_standard_rate(df: pd.DataFrame) -> pd.DataFrame:
    """
    Pay on the Standard formula: visits x revenue per visit x studio rate, kept
    between the $28 floor and the $83 maximum. Lets Starter classes be compared
    fairly, since all future classes are paid on Standard.
    """
    df = df.copy()
    rate = df["studio_num"].map(STUDIO_REV_SHARE)
    share_pay = (df["total_visits"] * df["revenue_per_visit"] * rate).round(2)
    df["formula_pay"] = share_pay.clip(lower=FLOOR_PAY)          # uncapped, for checks
    df["pay_if_standard_rate"] = share_pay.clip(FLOOR_PAY, MAX_PAY)
    df["hit_floor"] = share_pay < FLOOR_PAY
    return df


def validate(df: pd.DataFrame) -> None:
    # Only check pay that was actually recorded, not estimates
    real = df[~df["pay_estimated"]]
    std = real[real["pay_type"] == "Standard"]
    off = (std["total_class_pay"] - std["formula_pay"]).abs() > 0.01
    if off.any():
        raise ValueError(f"{off.sum()} Standard classes don't match the pay formula")
    starter = real[real["pay_type"] == "Starter"]
    if not (starter["total_class_pay"] == FLOOR_PAY).all():
        raise ValueError("A Starter class was paid something other than $28")
    print(f"Pay formula verified on all {len(std)} recorded Standard classes; "
          f"all {len(starter)} Starter classes paid $28")
    over = real[real["total_class_pay"] > MAX_PAY]
    for _, r in over.iterrows():
        print(f"NOTE: {r['class_date']:%b %d} {r['studio']} was paid "
              f"${r['total_class_pay']:.2f}, above the ${MAX_PAY:.0f} maximum")


EXCEL_COLUMNS = {  # column -> (description, number format)
    "class_date": ("Date the class ran", "yyyy-mm-dd"),
    "day_of_week": ("Day of the week", None),
    "time_label": ("Class start time", None),
    "studio": ("Studio 1 (capacity 10) or Studio 2 (capacity 16)", None),
    "slot": ("Studio + weekday + time, e.g. Studio 2 | Tue 5:30 PM", None),
    "shift_type": ("Regular (assigned schedule) or Picked up", None),
    "schedule_note": ("Holiday schedule or swapped shift, if a regular shift moved", None),
    "class_type": ("normal, or cancelled if no one attended", None),
    "total_visits": ("Clients who counted toward pay, including waitlist", "0"),
    "capacity": ("Spots in the room", "0"),
    "booked": ("Clients up to capacity", "0"),
    "waitlist": ("Clients over capacity", "0"),
    "fill_rate": ("Total visits / capacity (over 100% = waitlist)", "0%"),
    "is_full": ("TRUE if the class reached capacity", None),
    "pay_type": ("Starter ($28 flat) or Standard (revenue share), as actually paid", None),
    "total_class_pay": ("ACTUAL pay for the class ($28 for Starter; estimated where pay_estimated is TRUE)", "$#,##0.00"),
    "pay_estimated": ("TRUE if pay was estimated because there was no payslip", None),
    "hit_floor": ("TRUE if attendance was too low to beat the $28 minimum", None),
    "revenue_share_pct": ("Revenue share rate used (0% for Starter)", "0%"),
    "revenue_per_visit": ("Revenue per client visit", "$#,##0.00"),
    "month": ("Year-month", None),
    "pay_if_standard_rate": ("NOT what was paid. What the class would pay at the Standard rate "
                             "($28 min, $83 max); used only to compare slots fairly", "$#,##0.00"),
    "data_source": ("Payslip or Class log (logged by hand)", None),
    "data_note": ("Anything unusual about the row", None),
    "company": ("Anonymized company name", None),
    "coach": ("Anonymized coach name", None),
    "class_name": ("Class format", None),
    "source_file": ("Raw file the row came from", None),
}


def export_excel(out: pd.DataFrame, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    cols = [c for c in EXCEL_COLUMNS if c in out.columns]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        out[cols].to_excel(xw, sheet_name="Classes", index=False)
        ws = xw.sheets["Classes"]
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, start=1):
            cells = list(ws.iter_cols(min_col=i, max_col=i))[0]
            head = cells[0]
            head.font = Font(name="Arial", bold=True, color="FFFFFF")
            head.fill = PatternFill("solid", fgColor="3D5A80")
            head.alignment = Alignment(wrap_text=True, vertical="center")
            fmt = EXCEL_COLUMNS[col][1]
            for c in cells[1:]:
                c.font = Font(name="Arial", size=10)
                if fmt:
                    c.number_format = fmt
            width = max(len(str(c.value)) for c in cells if c.value is not None)
            ws.column_dimensions[head.column_letter].width = min(max(width + 2, 11), 38)

        dd = pd.DataFrame([(c, EXCEL_COLUMNS[c][0]) for c in cols],
                          columns=["Column", "Description"])
        dd.to_excel(xw, sheet_name="Data Dictionary", index=False)
        ws2 = xw.sheets["Data Dictionary"]
        ws2.column_dimensions["A"].width = 20
        ws2.column_dimensions["B"].width = 80
        for row in ws2.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=c.row == 1)


def main():
    df = pd.read_csv(IN_FILE, parse_dates=["class_date"])
    df["data_note"] = df["data_note"].fillna("")
    df["class_time"] = pd.to_datetime(df["class_time"], format="%H:%M:%S").dt.time

    df = add_time_fields(df)
    df = add_shift_type(df)
    df = add_attendance(df)
    df = add_pay_if_standard_rate(df)
    validate(df)

    counts = df["shift_type"].value_counts()
    print(f"Shift type: {counts.get('Regular', 0)} Regular, "
          f"{counts.get('Picked up', 0)} Picked up")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = df.copy()
    out["class_date"] = out["class_date"].dt.date
    out.to_csv(OUT_DIR / "classes_features.csv", index=False)
    export_excel(out, OUT_DIR / "classes_features.xlsx")
    print(f"Wrote {len(df)} rows to {OUT_DIR}")


if __name__ == "__main__":
    main()
